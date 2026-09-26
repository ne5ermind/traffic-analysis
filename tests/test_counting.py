from types import SimpleNamespace
import io
import pytest
from openpyxl import load_workbook
from ml.movements import aggregate, classify
from ml.artifacts import TrackStore
from backend.exports import xlsx_export, csv_export, cartogram_svg
from backend.validation import compare, parse_ground_truth


def track(tid=1, cls="car", t=0, y=0.5):
    return {
        "id": tid,
        "votes": {cls: 1},
        "confidence_sum": 4.5,
        "samples": 5,
        "first_seen": t,
        "last_seen": t + 4,
        "trajectory": [[0.1 + i * 0.2, y, t + i] for i in range(5)],
    }


def store_with_tracks(path):
    store = TrackStore(path)
    for i in range(5):
        store.save_track(track(i + 1, t=900 if i == 4 else 0, y=0.5 + i * 0.005))
    store.save_track(track(6, cls="pedestrian"))
    return store


def test_trajectory_classification_and_unknown(tmp_path):
    store = store_with_tracks(tmp_path / "tracks.db")
    definitions, _ = classify(store, {"zones": [], "lines": []})
    tracks = list(store.tracks())
    assert tracks[0]["movement_id"].startswith("auto_")
    assert tracks[-1]["movement_id"] == "unknown"  # one pedestrian is insufficient evidence
    result = aggregate(store, definitions, 1000)
    assert result["total_vehicles"] == 5 and result["pedestrians"] == 1
    assert result["intervals"][0]["total"] == 4 and result["intervals"][1]["total"] == 1
    assert result["diagnostics"]["unknown_tracks"] == 1
    assert sum(i["total"] for i in aggregate(store, definitions, 1000, 300)["intervals"]) == 5
    store.close()


def test_stationary_detection_is_not_counted_as_traffic(tmp_path):
    store = TrackStore(tmp_path / "tracks.db")
    stationary = track()
    stationary["trajectory"] = [[0.5 + (i % 2) * 0.001, 0.5, i] for i in range(5)]
    store.save_track(stationary)
    definitions, _ = classify(store, {"zones": [], "lines": []})
    result = aggregate(store, definitions, 10)
    saved = list(store.tracks())[0]
    store.close()
    assert not saved["countable"]
    assert result["total_vehicles"] == 0
    assert result["diagnostics"]["discarded_stationary_tracks"] == 1


def test_exact_interval_boundary_and_empty_interval(tmp_path):
    store = TrackStore(tmp_path / "tracks.db")
    store.save_track(track(t=896))
    result = aggregate(store, {}, 1800)
    assert result["intervals"][0]["total"] == 0
    assert result["intervals"][1]["total"] == 1
    assert len(result["intervals"]) == 2
    store.close()


def test_exports_and_formula_injection(tmp_path):
    db = tmp_path / "tracks.db"
    store = store_with_tracks(db)
    definitions, _ = classify(store, {"zones": [], "lines": []})
    result = aggregate(store, definitions, 1000)
    result.update(elapsed=2, model="test", device="cpu")
    result["movements"][0]["name"] = '=HYPERLINK("evil")'
    store.close()
    assert "'=HYPERLINK" in csv_export(result)
    svg = cartogram_svg(result)
    assert "<svg" in svg
    assert "Геометрия перекрёстка не определена" in svg
    assert 'marker-end="' not in svg
    p = SimpleNamespace(name="=evil", video={"width": 320}, result=result)
    out = io.BytesIO()
    xlsx_export(p, db, out)
    wb = load_workbook(io.BytesIO(out.getvalue()))
    assert wb.sheetnames == ["Summary", "Movements", "Time intervals", "Tracks", "Metadata"]
    assert wb["Summary"]["B2"].data_type == "s"
    assert wb["Summary"]["B2"].value == "'=evil"
    assert wb["Tracks"].max_row == 7


def test_validation_partial_zero_and_wrong_labels():
    result = {"movements": [{"id": "m1", "counts": {"car": 5, "truck": 2}}]}
    rows = parse_ground_truth(b"movement_id,class,count\nm1,car,0\n", "test.csv")
    report = compare(result, rows, 10)
    assert report["total"]["relative_error"] is None
    assert report["total"]["absolute_error"] == 5
    assert report["unlabelled_cells"] == 1
    with pytest.raises(ValueError):
        compare(result, [{"movement_id": "typo", "class": "car", "count": 5}], 10)
    with pytest.raises(ValueError):
        parse_ground_truth(b"movement_id,class,count\nm1,car,-1\n", "test.csv")
    with pytest.raises(ValueError):
        parse_ground_truth(b"movement_id,class,count\nm1,car,1\nm1,car,2", "test.csv")
