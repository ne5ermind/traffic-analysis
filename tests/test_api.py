import io
import json
import os
import tempfile
from pathlib import Path

import av

os.environ["INNOVA_DATA"] = tempfile.mkdtemp(prefix="innova-tests-")
from fastapi.testclient import TestClient  # noqa: E402
from app.main import app  # noqa: E402


def video_bytes(frames=10):
    file = io.BytesIO()
    with av.open(file, mode="w", format="mp4") as output:
        stream = output.add_stream("mpeg4", rate=5)
        stream.width = 32
        stream.height = 32
        for index in range(frames):
            from PIL import Image
            image = Image.new("RGB", (32, 32), (index * 15 % 256, 20, 30))
            frame = av.VideoFrame.from_image(image)
            for packet in stream.encode(frame):
                output.mux(packet)
        for packet in stream.encode():
            output.mux(packet)
    return file.getvalue()


def test_swagger_exposes_only_compact_public_api():
    with TestClient(app) as client:
        assert client.get("/docs").status_code == 200
        schema = client.get("/openapi.json").json()
    paths = {path for path in schema["paths"] if path.startswith("/api/")}
    assert schema["info"]["title"] == "InnovaTransport API"
    assert len(paths) == 10
    assert "/api/videos/{video_id}/media" in paths
    assert "/api/runs/{run_id}/results" in paths
    assert not paths & {
        "/api/videos/{video_id}/file", "/api/videos/{video_id}/frame",
        "/api/runs/{run_id}/aggregate", "/api/runs/{run_id}/cartogram",
        "/api/runs/{run_id}/detections", "/api/runs/{run_id}/cancel",
        "/api/runs/{run_id}/topology", "/api/events/{event_id}",
    }


def test_upload_scene_manual_review_export_and_reference():
    with TestClient(app) as client:
        assert client.get("/api/health").json() == {"ok": True}
        r = client.post("/api/videos", files={"file": ("test.mp4", video_bytes(), "video/mp4")})
        assert r.status_code == 200, r.text
        vid = r.json()["id"]
        assert client.get(f"/api/videos/{vid}/media?kind=frame&time_s=0.2").headers["content-type"] == "image/jpeg"
        scene = client.post("/api/scenes", json={"video_id": vid, "name": "S",
            "geometry": {"entries": [{"name": "N", "points": [[0.2, 0], [0.2, 1]]}],
                         "exits": [{"name": "S", "points": [[0.8, 0], [0.8, 1]]}], "roi": []}})
        assert scene.status_code == 200, scene.text
        assert client.post("/api/scenes", json={"video_id": vid, "name": "invalid",
            "geometry": {"entries": [], "exits": []}}).status_code == 422
        run = client.post("/api/runs", json={"scene_id": scene.json()["id"]})
        assert run.status_code == 200, run.text
        assert run.json()["tracker"] == "supervision-bytetrack"
        assert json.loads(run.json()["model"])["pipeline"] == "fast"
        assert json.loads(run.json()["model"])["models"] == ["models/yolo11s.pt"]
        run_id = run.json()["id"]
        event = client.post(f"/api/runs/{run_id}/events", json={"category": "car",
            "from_direction": "N", "to_direction": "S", "time_s": 0.5})
        assert event.status_code == 200, event.text
        event = event.json()
        changed = client.patch(f"/api/runs/{run_id}/events", json={"id": event["id"], "category": "bus",
            "from_direction": "N", "to_direction": "S", "time_s": 0.8,
            "status": "confirmed", "version": event["version"]})
        assert changed.status_code == 200, changed.text
        assert client.patch(f"/api/runs/{run_id}/events", json={"id": event["id"], "category": "bus",
            "from_direction": "N", "to_direction": "S", "time_s": 0.8,
            "status": "confirmed", "version": event["version"]}).status_code == 409
        assert client.get(f"/api/runs/{run_id}/results?view=summary").json()["matrix"] == [
            {"from": "N", "to": "S", "category": "bus", "count": 1}]
        summary = client.get(f"/api/runs/{run_id}/results?view=summary").json()
        assert summary["transport_total"] == 1 and summary["pedestrian_total"] == 0
        assert summary["category_totals"] == [{"category": "bus", "count": 1}]
        assert summary["video"]["name"] == "test.mp4" and summary["scene"]["name"] == "S"
        from openpyxl import load_workbook
        exported = client.get(f"/api/runs/{run_id}/export/xlsx")
        assert exported.status_code == 200
        book = load_workbook(io.BytesIO(exported.content))
        assert book.sheetnames[:2] == ["Отчёт", "Подсчёты"]
        assert book["Отчёт"]["E6"].value == 1
        counts = book["Подсчёты"]
        assert counts["E1"].value == "Направление 1\nN → S"
        assert counts["E2"].value == "ОТ" and counts["E3"].value == 1
        assert counts["A1"].fill.fgColor.rgb.endswith("0B1F3A")
        assert book["Flow Matrix"]["D2"].value == 1
        assert book["Cartogram coverage"]["A2"].value == "topology"
        assert client.get(f"/api/runs/{run_id}/export/json").json()["cartogram"]["total"] == 1
        gt = client.post(f"/api/videos/{vid}/ground-truth", files={"file": (
            "reference.csv", b"from,to,category,count\nN,S,bus,1\n", "text/csv")})
        assert gt.status_code == 200, gt.text
        evaluation = client.get(f"/api/runs/{run_id}/evaluate/{gt.json()['id']}")
        assert evaluation.json()["wape"] == 0


def test_video_paths_survive_docker_prefix_changes_and_keep_media_type():
    from app.storage import connection

    with TestClient(app) as client:
        uploaded = client.post("/api/videos", files={"file": (
            "portable.mov", video_bytes(), "video/quicktime")})
        assert uploaded.status_code == 200, uploaded.text
        video = uploaded.json()
        assert video["path"].startswith("uploads/")
        filename = Path(video["path"]).name
        with connection() as db:
            db.execute("UPDATE videos SET path=? WHERE id=?",
                       (f"/srv/innova/data/uploads/{filename}", video["id"]))
        response = client.get(f"/api/videos/{video['id']}/media?kind=file")
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("video/quicktime")
        assert client.get(f"/api/videos/{video['id']}/media?kind=frame&time_s=0").status_code == 200
        listed = next(item for item in client.get("/api/videos").json() if item["id"] == video["id"])
        assert listed["available"] is True
        with connection() as db:
            db.execute("UPDATE videos SET path=? WHERE id=?",
                       ("/srv/innova/data/uploads/not-present.mov", video["id"]))
        listed = next(item for item in client.get("/api/videos").json() if item["id"] == video["id"])
        assert listed["available"] is False
        assert client.get(f"/api/videos/{video['id']}/media?kind=file").status_code == 404


def test_model_resolver_rebinds_legacy_prefix_but_rejects_external_files(tmp_path):
    from app.storage import resolve_model_path

    legacy = resolve_model_path("/srv/innova/data/models/yolo11s.pt")
    assert legacy.is_file() and legacy.name == "yolo11s.pt"
    external = tmp_path / "external.pt"
    external.write_bytes(b"not a model")
    resolved = resolve_model_path(external)
    assert resolved != external.resolve()
    assert resolved.name == "external.pt" and not resolved.is_file()


def test_automatic_run_counts_and_saved_boxes_without_manual_lines(monkeypatch):
    from app.worker import run as process
    from app.storage import connection

    class Scalar:
        def __init__(self, value):
            self.value = value

        def item(self):
            return self.value

    class Coordinates:
        def __init__(self, offset=0):
            self.offset = offset

        def tolist(self):
            return [2 + self.offset, 3, 15 + self.offset, 27]

    class Box:
        def __init__(self, identity, category, offset=0):
            self.id = Scalar(identity)
            self.cls = Scalar(category)
            self.conf = Scalar(0.9)
            self.xyxy = [Coordinates(offset)]

    class Boxes(list):
        id = True

    class Model:
        def __init__(self):
            self.frames = 0

        def track(self, *args, **kwargs):
            self.frames += 1
            offset = min(self.frames - 1, 3)
            boxes = Boxes([Box(7, 2, offset), Box(8, 0, offset)])
            if self.frames == 1:
                boxes.append(Box(9, 3))  # One-frame false positive is not counted.
            return [type("Result", (), {"boxes": boxes})()]

    monkeypatch.setattr("app.model.load_yolo", lambda path: Model())
    with TestClient(app) as client:
        video = client.post("/api/videos", files={"file": ("auto.mp4", video_bytes(), "video/mp4")}).json()
        assert client.post("/api/scenes", json={"video_id": video["id"], "name": "wrong",
            "geometry": {"mode": "automatic", "entries": [{"name": "X", "points": [[0, 0], [1, 1]]}],
                         "exits": []}}).status_code == 422
        created = client.post("/api/runs", json={"video_id": video["id"],
            "model_secondary": None, "tracker": "bytetrack.yaml"})
        assert created.status_code == 200, created.text
        run_id = created.json()["id"]
        assert client.post("/api/runs", json={"scene_id": created.json()["scene_id"],
            "video_id": "other"}).status_code == 422
        assert client.get(f"/api/runs/{run_id}/results?view=detections&time_s=0").json()["objects"] == []
        with connection() as db:
            job = dict(db.execute("SELECT * FROM runs WHERE id=?", (run_id,)).fetchone())
        process(job)
        summary = client.get(f"/api/runs/{run_id}/results?view=summary").json()
        assert summary["total"] == 2
        assert summary["geometry"]["mode"] == "automatic"
        assert {row["category"]: row["count"] for row in summary["matrix"]} == {"car": 1, "pedestrian": 1}
        events = client.get(f"/api/runs/{run_id}/results?view=events").json()
        assert {event["tracker_id"] for event in events} == {7, 8}
        assert all(event["status"] == "auto_accept" for event in events)
        frame = client.get(f"/api/runs/{run_id}/results?view=detections&time_s=0.6").json()
        assert {obj["category"] for obj in frame["objects"]} == {"car", "pedestrian"}
        assert all(all(0 <= coord <= 1 for coord in obj["box"]) for obj in frame["objects"])
        assert client.get(f"/api/runs/{run_id}/results?view=detections&time_s=10").json()["objects"] == []
        assert client.post(f"/api/runs/{run_id}/events", json={"category": "car",
            "from_direction": "N", "to_direction": "S", "time_s": 0.2}).status_code == 422
        assert client.post(f"/api/runs/{run_id}/events", json={"category": "car",
            "from_direction": "В кадре", "to_direction": "В кадре", "time_s": 0.2}).status_code == 200
        assert client.get(f"/api/runs/{run_id}/results?view=summary").json()["total"] == 3
        edited_scene = client.put(f"/api/scenes/{created.json()['scene_id']}", json={
            "video_id": video["id"], "name": "Новая разметка", "geometry": {
                "mode": "manual", "entries": [{"name": "N", "points": [[0.2, 0], [0.2, 1]]}],
                "exits": [{"name": "S", "points": [[0.8, 0], [0.8, 1]]}], "roi": []}})
        assert edited_scene.status_code == 200, edited_scene.text
        assert client.get(f"/api/runs/{run_id}/results?view=summary").json()["geometry"]["mode"] == "automatic"
        assert client.get(f"/api/runs/{run_id}/results?view=summary").json()["total"] == 3


def test_ensemble_worker_integrates_tracked_boxes_with_events(monkeypatch):
    import numpy as np
    import supervision as sv
    from app.worker import run as process
    from app.storage import connection

    calls = []

    class Processor:
        def __init__(self, primary, secondary, fps):
            calls.append((Path(primary).name, Path(secondary).name, fps))
            self.frame = 0

        def process_frame(self, frame):
            calls.append("frame")
            offset = min(self.frame, 3)
            self.frame += 1
            detections = sv.Detections(
                xyxy=np.array([[2 + offset, 3, 15 + offset, 27],
                               [18 + offset, 4, 28 + offset, 29]], dtype=np.float32),
                confidence=np.array([0.8, 0.9], dtype=np.float32),
                class_id=np.array([2, 0], dtype=np.int64))
            detections.tracker_id = np.array([42, 43], dtype=np.int64)
            return detections

    monkeypatch.setattr("app.ensemble.EnsembleFrameProcessor", Processor)
    with TestClient(app) as client:
        video = client.post("/api/videos", files={"file": ("ensemble.mp4", video_bytes(), "video/mp4")}).json()
        created = client.post("/api/runs", json={"video_id": video["id"], "profile": "ensemble"})
        assert created.status_code == 200, created.text
        run_id = created.json()["id"]
        with connection() as db:
            job = dict(db.execute("SELECT * FROM runs WHERE id=?", (run_id,)).fetchone())
        process(job)
        assert calls[0] == ("yolo11m.pt", "yolov10m.pt", 5.0)
        assert calls.count("frame") == 10
        summary = client.get(f"/api/runs/{run_id}/results?view=summary").json()
        assert summary["total"] == 2
        assert {event["tracker_id"] for event in client.get(f"/api/runs/{run_id}/results?view=events").json()} == {42, 43}
        objects = client.get(f"/api/runs/{run_id}/results?view=detections&time_s=0.6").json()["objects"]
        assert {(obj["tracker_id"], obj["class_id"]) for obj in objects} == {(42, 2), (43, 0)}
        assert objects[0]["xyxy"] == [5, 3, 18, 27]
        assert objects[0]["box"] == [0.15625, 0.09375, 0.5625, 0.84375]


def test_ensemble_track_id_flows_into_manual_directions(monkeypatch):
    import numpy as np
    import supervision as sv
    from app.worker import run as process
    from app.storage import connection

    class Processor:
        def __init__(self, primary, secondary, fps):
            self.frame = 0

        def process_frame(self, image):
            center = [0.1, 0.4, 0.9, 0.95][min(self.frame, 3)] * 32
            self.frame += 1
            tracked = sv.Detections(
                xyxy=np.array([[center - 1, 3, center + 1, 28]], dtype=np.float32),
                confidence=np.array([0.9], dtype=np.float32),
                class_id=np.array([2], dtype=np.int64))
            tracked.tracker_id = np.array([52], dtype=np.int64)
            return tracked

    monkeypatch.setattr("app.ensemble.EnsembleFrameProcessor", Processor)
    with TestClient(app) as client:
        video = client.post("/api/videos", files={"file": ("routes.mp4", video_bytes(), "video/mp4")}).json()
        scene = client.post("/api/scenes", json={"video_id": video["id"], "name": "Маршрут",
            "geometry": {"entries": [{"name": "Север", "points": [[0.2, 0], [0.2, 1]]}],
                         "exits": [{"name": "Юг", "points": [[0.8, 0], [0.8, 1]]}], "roi": []}}).json()
        created = client.post("/api/runs", json={"scene_id": scene["id"], "profile": "ensemble"}).json()
        with connection() as db:
            job = dict(db.execute("SELECT * FROM runs WHERE id=?", (created["id"],)).fetchone())
        process(job)
        events = client.get(f"/api/runs/{created['id']}/results?view=events").json()
        assert len(events) == 1
        assert (events[0]["tracker_id"], events[0]["from_direction"], events[0]["to_direction"]) == (
            52, "Север", "Юг")
        assert events[0]["status"] == "review"
        assert client.patch(f"/api/runs/{created['id']}",
            json={"action": "set_topology", "topology": "three_way"}).status_code == 200
        confirmed = client.patch(f"/api/runs/{created['id']}/events", json={"id": events[0]["id"],
            "category": "car", "from_direction": "Север", "to_direction": "Юг",
            "time_s": events[0]["time_s"], "status": "confirmed", "version": events[0]["version"]})
        assert confirmed.status_code == 200, confirmed.text
        diagram = client.get(f"/api/runs/{created['id']}/results?view=cartogram").json()
        assert diagram["topology"] == "three_way"
        assert diagram["routed"] == 1 and diagram["total"] == 1


def test_ensemble_rejects_mismatched_configuration():
    with TestClient(app) as client:
        video = client.post("/api/videos", files={"file": ("params.mp4", video_bytes(), "video/mp4")}).json()
        vid = video["id"]
        for payload in (
            {"profile": "fast", "model_secondary": "data/models/yolov10m.pt"},
            {"profile": "ensemble", "tracker": "botsort.yaml"},
            {"profile": "ensemble", "model_secondary": "data/models/yolo11m.pt"},
            {"profile": "ensemble", "model_secondary": "../other.pt"},
            {"sample_fps": 30},
        ):
            assert client.post("/api/runs", json={"video_id": vid, **payload}).status_code == 422
        alternate = client.post("/api/runs", json={"video_id": vid,
            "profile": "ensemble", "model": "data/models/yolov8m.pt",
            "model_secondary": "data/models/yolov10m.pt"})
        assert alternate.status_code == 200, alternate.text
        assert [Path(path).name for path in json.loads(alternate.json()["model"])["models"]] == [
            "yolov8m.pt", "yolov10m.pt"]
        legacy_client = client.post("/api/runs", json={"video_id": vid,
            "model": "data/models/yolov8m.pt", "model_secondary": "data/models/yolov10m.pt"})
        assert legacy_client.status_code == 200
        assert json.loads(legacy_client.json()["model"])["fusion"] == "wbf"


def test_fast_worker_samples_before_inference_and_saves_real_timestamps(monkeypatch):
    import numpy as np
    import supervision as sv
    from app.storage import connection
    from app.worker import run as process

    calls = []

    class Processor:
        def __init__(self, weights, fps, image_size):
            calls.append((Path(weights).name, fps, image_size))
            self.frame = 0

        def process_frame(self, frame):
            calls.append("inference")
            offset = min(self.frame, 3)
            self.frame += 1
            result = sv.Detections(
                xyxy=np.array([[2 + offset, 3, 15 + offset, 27]], dtype=np.float32),
                confidence=np.array([0.9], dtype=np.float32),
                class_id=np.array([2], dtype=np.int64))
            result.tracker_id = np.array([71], dtype=np.int64)
            return result

    monkeypatch.setattr("app.ensemble.FastFrameProcessor", Processor)
    with TestClient(app) as client:
        video = client.post("/api/videos", files={"file": ("fast.mp4", video_bytes(), "video/mp4")}).json()
        created = client.post("/api/runs", json={"video_id": video["id"], "sample_fps": 2}).json()
        with connection() as db:
            job = dict(db.execute("SELECT * FROM runs WHERE id=?", (created["id"],)).fetchone())
        process(job)
        assert calls[0] == ("yolo11s.pt", 2, 640)
        assert 3 <= calls.count("inference") <= 5  # Source: 10 frames / 2 seconds at 5 fps.
        events = client.get(f"/api/runs/{created['id']}/results?view=events").json()
        assert len(events) == 1 and events[0]["tracker_id"] == 71
        result = client.get(f"/api/runs/{created['id']}").json()
        assert result["status"] == "completed" and result["progress"] == 1
        frame = client.get(f"/api/runs/{created['id']}/results?view=detections&time_s=1.5").json()
        assert frame["objects"][0]["tracker_id"] == 71
        assert client.get(f"/api/runs/{created['id']}/results?view=detections&time_s=1.7").json()["objects"]


def test_fast_worker_reduces_frequency_when_runtime_budget_is_at_risk(monkeypatch):
    import supervision as sv
    from app.storage import connection
    from app.worker import run as process

    calls = []

    class Processor:
        def __init__(self, weights, fps, image_size):
            pass

        def process_frame(self, frame):
            calls.append(1)
            return sv.Detections.empty()

    monkeypatch.setattr("app.ensemble.FastFrameProcessor", Processor)
    monkeypatch.setattr("app.worker.budgeted_fps", lambda *args: 2.0)
    with TestClient(app) as client:
        video = client.post("/api/videos", files={"file": (
            "long.mp4", video_bytes(frames=610), "video/mp4")}).json()
        created = client.post("/api/runs", json={"video_id": video["id"]}).json()
        with connection() as db:
            job = dict(db.execute("SELECT * FROM runs WHERE id=?", (created["id"],)).fetchone())
        process(job)
        run = client.get(f"/api/runs/{created['id']}").json()
        assert run["status"] == "completed" and "частота: до 2 кадр/с" in run["message"]
        assert 280 < len(calls) < 450  # First 30 seconds at source fps 5, remainder at 2.


def test_stalled_worker_is_reported_and_can_be_cancelled():
    from app.storage import connection

    with TestClient(app) as client:
        video = client.post("/api/videos", files={"file": ("stalled.mp4", video_bytes(), "video/mp4")}).json()
        run_id = client.post("/api/runs", json={"video_id": video["id"]}).json()["id"]
        with connection() as db:
            db.execute("UPDATE runs SET status='running',heartbeat=? WHERE id=?",
                       ("2020-01-01T00:00:00+00:00", run_id))
        detail = client.get(f"/api/runs/{run_id}").json()
        assert detail["status"] == "stalled" and "Worker" in detail["message"]
        assert client.get(f"/api/runs?video_id={video['id']}").json()[0]["status"] == "stalled"
        assert client.patch(f"/api/runs/{run_id}", json={"action": "cancel"}).json()["status"] == "cancelled"


def test_supplied_example_four_starts_with_triangular_cartogram():
    with TestClient(app) as client:
        video = client.post("/api/videos", files={"file": (
            "Пример 4_утро_1.mov", video_bytes(), "video/quicktime")}).json()
        created = client.post("/api/runs", json={"video_id": video["id"]}).json()
        run_id = created["id"]
        geometry = client.get(f"/api/runs/{run_id}").json()["geometry"]
        assert geometry["topology"] == "three_way" and geometry["topology_source"] == "reference"
        assert client.get(f"/api/runs/{run_id}/results?view=cartogram").json()["topology"] == "three_way"
        assert client.patch(f"/api/runs/{run_id}",
                            json={"action": "set_topology", "topology": "four_way"}).status_code == 200
        assert client.get(f"/api/runs/{run_id}").json()["geometry"]["topology_source"] == "user"
