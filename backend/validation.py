import csv
import io
import math
from collections import defaultdict
from openpyxl import load_workbook
from ml.detectors.base import CLASSES


def parse_ground_truth(content, filename):
    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.DictReader(io.StringIO(text), dialect=dialect))
    elif filename.lower().endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = wb["Ground truth"] if "Ground truth" in wb.sheetnames else wb.active
            iterator = sheet.iter_rows(values_only=True)
            header = next(iterator)
            rows = [dict(zip(header, row)) for row in iterator if any(v is not None for v in row)]
        finally:
            wb.close()
    else:
        raise ValueError("Загрузите CSV или XLSX")
    if not rows or len(rows) > 100000:
        raise ValueError("Эталон должен содержать от 1 до 100000 строк")
    normalized = []
    seen = set()
    for row in rows:
        if row.get("count") is None or str(row.get("count")).strip() == "":
            continue  # Empty template cells are explicitly unlabelled, not zero.
        mid, cls = str(row.get("movement_id", "")).strip(), str(row.get("class", "")).strip()
        if not mid or cls not in CLASSES:
            raise ValueError("Нужны колонки movement_id, class, count; class: " + ", ".join(CLASSES))
        try:
            count = float(row["count"])
        except (ValueError, TypeError, KeyError) as e:
            raise ValueError("count должен быть целым неотрицательным числом") from e
        if not math.isfinite(count) or count < 0 or count != int(count):
            raise ValueError("count должен быть целым неотрицательным числом")
        if (mid, cls) in seen:
            raise ValueError("Повтор пары movement_id + class в эталоне")
        seen.add((mid, cls))
        normalized.append(dict(movement_id=mid, **{"class": cls}, count=int(count)))
    if not normalized:
        raise ValueError("В эталоне нет заполненных значений count")
    return normalized


def compare(result, truth, threshold):
    automatic = {(m["id"], c): n for m in result["movements"] for c, n in m["counts"].items()}
    expected = {(r["movement_id"], r["class"]): r["count"] for r in truth}
    known = {m["id"] for m in result["movements"]} | {"unknown"}
    unknown_ids = sorted({k[0] for k in expected} - known)
    if unknown_ids:
        raise ValueError("В эталоне неизвестные movement_id: " + ", ".join(unknown_ids))

    def metrics(manual, auto):
        error = abs(auto - manual)
        relative = error / manual * 100 if manual else (0 if auto == 0 else None)
        return dict(
            manual=manual,
            automatic=auto,
            absolute_error=error,
            relative_error=relative,
            within_threshold=error == 0 if relative is None else relative <= threshold,
        )

    rows, by_class, by_movement = [], defaultdict(lambda: [0, 0]), defaultdict(lambda: [0, 0])
    # Explicit scope: omitted cells are unlabelled, never assumed to be zero.
    for (mid, cls), manual in expected.items():
        auto = automatic.get((mid, cls), 0)
        rows.append(dict(movement_id=mid, **{"class": cls}, **metrics(manual, auto)))
        for group, key in ((by_class, cls), (by_movement, mid)):
            group[key][0] += manual
            group[key][1] += auto
    totals = metrics(sum(r["manual"] for r in rows), sum(r["automatic"] for r in rows))
    totals["sum_absolute_error"] = sum(r["absolute_error"] for r in rows)
    totals["weighted_absolute_percentage_error"] = (
        totals["sum_absolute_error"] / totals["manual"] * 100 if totals["manual"] else (0 if totals["sum_absolute_error"] == 0 else None)
    )
    return dict(
        rows=rows,
        by_class={k: metrics(*v) for k, v in by_class.items()},
        by_movement={k: metrics(*v) for k, v in by_movement.items()},
        total=totals,
        threshold=threshold,
        compared_cells=len(rows),
        unlabelled_cells=len(set(automatic) - set(expected)),
        scope="Сравниваются только явно заданные пары направление + класс",
    )
