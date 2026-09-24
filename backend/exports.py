import csv
import io
import json
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from backend.cartogram import cartogram_svg  # noqa: F401
from ml.artifacts import TrackStore
from ml.detectors.base import CLASSES

HEADERS = ["movement_id", "movement", *CLASSES, "total_vehicles", "reliability"]


def safe_cell(value):
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def movement_rows(result):
    yield HEADERS
    for m in result["movements"]:
        yield [m["id"], m["name"], *[m["counts"][c] for c in CLASSES], m["total"], m["reliability"]]


def csv_export(result):
    out = io.StringIO()
    out.write("\ufeff")
    writer = csv.writer(out)
    for row in movement_rows(result):
        writer.writerow([safe_cell(v) for v in row])
    return out.getvalue()


def xlsx_export(project, db_path, output):
    wb = Workbook(write_only=True)

    def sheet(name, rows):
        ws = wb.create_sheet(name)
        ws.freeze_panes = "A2"
        for col in "ABCDEFGHIJK":
            ws.column_dimensions[col].width = 22
        for i, row in enumerate(rows):
            cells = []
            for value in row:
                cell = WriteOnlyCell(ws, value=safe_cell(value))
                if i == 0:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="334155")
                cells.append(cell)
            ws.append(cells)

    result = project.result
    sheet(
        "Summary",
        [
            ["Показатель", "Значение"],
            ["Проект", project.name],
            ["Всего транспорта", result["total_vehicles"]],
            ["Пешеходы", result["pedestrians"]],
            *[[c, n] for c, n in result["counts"].items()],
            ["Время анализа, с", result["elapsed"]],
        ],
    )
    sheet("Movements", movement_rows(result))
    sheet(
        "Time intervals",
        [
            ["start_seconds", "end_seconds", *CLASSES, "total_vehicles"],
            *[[r["start"], r["end"], *[r[c] for c in CLASSES], r["total"]] for r in result["intervals"]],
        ],
    )
    store = TrackStore(db_path)
    try:

        def rows():
            yield [
                "track_id",
                "class",
                "confidence",
                "first_seen",
                "last_seen",
                "entry_zone",
                "exit_zone",
                "movement_id",
                "counted_at",
                "countable",
                "trajectory",
            ]
            for t in store.tracks():
                yield [
                    t["id"],
                    t["class_name"],
                    t["confidence"],
                    t["first_seen"],
                    t["last_seen"],
                    t["entry_zone"],
                    t["exit_zone"],
                    t["movement_id"],
                    t["counted_at"],
                    bool(t["countable"]),
                    json.dumps(t["trajectory"], separators=(",", ":"))[:32767],
                ]

        sheet("Tracks", rows())
    finally:
        store.close()
    sheet(
        "Metadata",
        [
            ["Параметр", "Значение"],
            *[
                [k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v]
                for k, v in {
                    **project.video,
                    "model": result.get("model"),
                    "device": result.get("device"),
                    "calibration": result.get("calibration"),
                }.items()
            ],
        ],
    )
    wb.save(output)
