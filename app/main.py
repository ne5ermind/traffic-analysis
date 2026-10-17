"""Local InnovaTransport HTTP service."""
import csv
import io
import json
import re
from bisect import bisect_right
from collections import Counter, defaultdict
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Literal
from uuid import uuid4

from fastapi import FastAPI, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field, field_validator, model_validator

from .storage import (ROOT, connection, init, parse_event, portable_data_path,
                      portable_model_path, resolve_data_path, rows)

STATIC = Path(__file__).resolve().parents[1] / "static"
ALLOWED_VIDEO = {".mp4", ".mov", ".mkv", ".avi", ".webm"}
VIDEO_MEDIA_TYPES = {
    ".mp4": "video/mp4",
    ".mov": "video/quicktime",
    ".mkv": "video/x-matroska",
    ".avi": "video/x-msvideo",
    ".webm": "video/webm",
}
STATUSES = {"review", "confirmed", "auto_accept", "discarded"}


@asynccontextmanager
async def lifespan(_app):
    init()
    yield


app = FastAPI(
    title="InnovaTransport API",
    version="1.1.0",
    description=(
        "Компактный API для загрузки видео, запуска анализа и просмотра "
        "результатов. Интерактивная Swagger-документация доступна по адресу `/docs`."
    ),
    docs_url="/docs",
    redoc_url=None,
    openapi_url="/openapi.json",
    openapi_tags=[
        {"name": "Система", "description": "Проверка готовности сервиса."},
        {"name": "Видео", "description": "Загрузка, каталог и медиаданные."},
        {"name": "Зоны", "description": "Зоны и линии подсчёта."},
        {"name": "Анализ", "description": "Запуск, статус и управление анализом."},
        {"name": "Результаты", "description": "Сводка, события, картограмма и экспорт."},
    ],
    lifespan=lifespan,
)
app.mount("/static", StaticFiles(directory=STATIC), name="static")


def now():
    return datetime.now(timezone.utc).isoformat()


def one(db, table, identity):
    if table not in ("videos", "scenes", "runs", "events", "ground_truth"):
        raise ValueError("unknown table")
    row = db.execute(f"SELECT * FROM {table} WHERE id=?", (identity,)).fetchone()
    if not row:
        raise HTTPException(404, f"{table}: не найдено")
    return dict(row)


def geometry_check(geometry):
    if geometry.get("mode") == "automatic":
        if geometry.get("entries") or geometry.get("exits") or geometry.get("roi"):
            raise ValueError("Автоматический режим не использует линии и области")
        return
    if geometry.get("mode", "manual") != "manual":
        raise ValueError("Неизвестный режим подсчёта")
    if not geometry.get("entries") or not geometry.get("exits"):
        raise ValueError("Нужна хотя бы одна линия въезда и выезда")
    for kind in ("entries", "exits"):
        names = set()
        for line in geometry[kind]:
            if not line.get("name") or line["name"] in names:
                raise ValueError("У каждой линии должно быть уникальное название")
            names.add(line["name"])
            if len(line.get("points", [])) != 2:
                raise ValueError("Линия состоит из двух точек")
            for point in line["points"]:
                if len(point) != 2 or not all(0 <= float(v) <= 1 for v in point):
                    raise ValueError("Координаты должны быть нормализованы (0..1)")
    for point in geometry.get("roi", []):
        if len(point) != 2 or not all(0 <= float(v) <= 1 for v in point):
            raise ValueError("Некорректная область интереса")


class SceneIn(BaseModel):
    video_id: str
    name: str = Field(min_length=1, max_length=120)
    geometry: dict


class RunIn(BaseModel):
    scene_id: str | None = None
    video_id: str | None = None
    model: str | None = None
    model_secondary: str | None = None
    profile: Literal["fast", "ensemble"] = "fast"
    sample_fps: float = Field(default=6, ge=2, le=12)
    topology: Literal["auto", "three_way", "four_way"] = "auto"
    tracker: str = "supervision-bytetrack"

    @field_validator("tracker")
    @classmethod
    def tracker_allowed(cls, v):
        if v not in ("supervision-bytetrack", "bytetrack.yaml", "botsort.yaml"):
            raise ValueError("Неизвестный трекер")
        return v

    @model_validator(mode="after")
    def tracker_matches_models(self):
        if (self.model_secondary is not None and "profile" not in self.model_fields_set
                and self.tracker == "supervision-bytetrack"):
            # Clients of the original API supplied the second model without a profile.
            self.profile = "ensemble"
        if self.tracker != "supervision-bytetrack" and (self.profile != "fast" or self.model_secondary):
            raise ValueError("Старые трекеры поддерживают только один YOLO")
        if self.tracker == "supervision-bytetrack" and self.profile == "fast" and self.model_secondary:
            raise ValueError("Быстрый режим использует одну модель; выберите режим ансамбля")
        return self


class TopologyIn(BaseModel):
    topology: Literal["auto", "three_way", "four_way"]


class RunUpdate(BaseModel):
    action: Literal["cancel", "set_topology"]
    topology: Literal["auto", "three_way", "four_way"] | None = None

    @model_validator(mode="after")
    def topology_required(self):
        if self.action == "set_topology" and self.topology is None:
            raise ValueError("Для set_topology нужно указать topology")
        return self


class EventIn(BaseModel):
    category: str
    from_direction: str
    to_direction: str
    time_s: float = Field(ge=0)
    start_s: float | None = None
    status: str = "confirmed"
    reason: str = ""

    @field_validator("status")
    @classmethod
    def status_allowed(cls, v):
        if v not in STATUSES:
            raise ValueError("Недопустимый статус")
        return v


class EventEdit(EventIn):
    id: str
    version: int = Field(ge=1)


@app.get("/", include_in_schema=False)
def home():
    return FileResponse(STATIC / "index.html")


@app.get("/api/health", tags=["Система"], summary="Проверить готовность API")
def health():
    return {"ok": True}


@app.post("/api/videos", tags=["Видео"], summary="Загрузить видео")
async def upload(file: UploadFile = File(...)):
    extension = Path(file.filename or "").suffix.lower()
    if extension not in ALLOWED_VIDEO:
        raise HTTPException(422, "Поддерживаются MP4, MOV, MKV, AVI и WebM")
    identity = str(uuid4())
    folder = ROOT / "uploads"
    folder.mkdir(exist_ok=True)
    destination = folder / (identity + extension)
    with destination.open("wb") as target:
        while chunk := await file.read(1024 * 1024):
            target.write(chunk)
    if destination.stat().st_size == 0:
        destination.unlink()
        raise HTTPException(422, "Пустой файл")
    try:
        import av
        with av.open(str(destination)) as source:
            stream = source.streams.video[0]
            width, height = stream.width, stream.height
            duration = (float(stream.duration * stream.time_base) if stream.duration is not None
                        else float(source.duration or 0) / 1_000_000)
            fps = float(stream.average_rate or 25)
            if duration <= 0 and stream.frames:
                duration = stream.frames / fps
    except Exception as exc:
        destination.unlink(missing_ok=True)
        raise HTTPException(422, f"Невозможно прочесть видео: {exc}") from exc
    with connection() as db:
        db.execute("INSERT INTO videos VALUES(?,?,?,?,?,?,?,?)",
                   (identity, file.filename, portable_data_path(destination),
                    width, height, duration, fps, now()))
        return one(db, "videos", identity)


@app.get("/api/videos", tags=["Видео"], summary="Получить каталог видео")
def videos():
    with connection() as db:
        result = rows(db, "SELECT id,name,path,width,height,duration,fps,created_at FROM videos ORDER BY created_at DESC")
    for video in result:
        video["available"] = resolve_data_path(video.pop("path")).is_file()
    return result


def video_file(video_id: str):
    with connection() as db:
        video = one(db, "videos", video_id)
    path = resolve_data_path(video["path"])
    if not path.is_file():
        raise HTTPException(404, "Файл видеозаписи не найден в хранилище")
    return FileResponse(path, media_type=VIDEO_MEDIA_TYPES.get(path.suffix.lower(), "application/octet-stream"))


def frame(video_id: str, time_s: float = Query(0, ge=0)):
    import av
    from PIL import Image
    with connection() as db:
        video = one(db, "videos", video_id)
    path = resolve_data_path(video["path"])
    if not path.is_file():
        raise HTTPException(404, "Файл видеозаписи не найден в хранилище")
    with av.open(str(path)) as source:
        stream = source.streams.video[0]
        try:
            source.seek(int(time_s / float(stream.time_base)), stream=stream, backward=True)
        except Exception:
            source.seek(0)
        found = None
        for item in source.decode(stream):
            found = item
            if item.pts is None or float(item.pts * item.time_base) >= time_s:
                break
        if found is None:
            raise HTTPException(422, "Нет кадров")
        result = io.BytesIO()
        Image.fromarray(found.to_ndarray(format="rgb24")).save(result, format="JPEG", quality=85)
        return Response(result.getvalue(), media_type="image/jpeg")


@app.get(
    "/api/videos/{video_id}/media",
    tags=["Видео"],
    summary="Получить видео или preview-кадр",
    description="`kind=file` возвращает исходное видео, `kind=frame` — JPEG-кадр на моменте `time_s`.",
)
def video_media(
    video_id: str,
    kind: Literal["file", "frame"] = "file",
    time_s: float = Query(0, ge=0),
):
    return frame(video_id, time_s) if kind == "frame" else video_file(video_id)


@app.post("/api/scenes", tags=["Зоны"], summary="Создать зону подсчёта")
def create_scene(payload: SceneIn):
    try:
        geometry_check(payload.geometry)
    except (ValueError, KeyError, TypeError) as exc:
        raise HTTPException(422, str(exc)) from exc
    identity = str(uuid4())
    with connection() as db:
        one(db, "videos", payload.video_id)
        db.execute("INSERT INTO scenes VALUES(?,?,?,?,?)",
                   (identity, payload.video_id, payload.name, json.dumps(payload.geometry), now()))
        return one(db, "scenes", identity)


@app.get("/api/scenes", tags=["Зоны"], summary="Получить зоны подсчёта")
def scenes(video_id: str | None = None):
    with connection() as db:
        output = rows(db, "SELECT * FROM scenes WHERE video_id=? ORDER BY created_at DESC", (video_id,)) if video_id else rows(db, "SELECT * FROM scenes ORDER BY created_at DESC")
    for scene in output:
        scene["geometry"] = json.loads(scene["geometry"])
    return output


@app.put("/api/scenes/{scene_id}", tags=["Зоны"], summary="Изменить зону подсчёта")
def update_scene(scene_id: str, payload: SceneIn):
    try:
        geometry_check(payload.geometry)
    except (ValueError, KeyError, TypeError) as exc:
        raise HTTPException(422, str(exc)) from exc
    with connection() as db:
        old = one(db, "scenes", scene_id)
        if old["video_id"] != payload.video_id:
            raise HTTPException(409, "Запись сцены нельзя привязать к другому видео")
        db.execute("UPDATE scenes SET name=?,geometry=? WHERE id=?",
                   (payload.name, json.dumps(payload.geometry), scene_id))
        return {**one(db, "scenes", scene_id), "geometry": payload.geometry}


@app.post("/api/runs", tags=["Анализ"], summary="Запустить анализ")
def create_run(payload: RunIn):
    identity = str(uuid4())
    with connection() as db:
        if payload.scene_id:
            scene = one(db, "scenes", payload.scene_id)
            if payload.video_id and scene["video_id"] != payload.video_id:
                raise HTTPException(422, "Сцена относится к другому видео")
        elif payload.video_id:
            one(db, "videos", payload.video_id)
            automatic = json.dumps({"mode": "automatic", "entries": [], "exits": [], "roi": []})
            scene = db.execute("SELECT * FROM scenes WHERE video_id=? AND geometry=? LIMIT 1",
                               (payload.video_id, automatic)).fetchone()
            if scene is None:
                scene_id = str(uuid4())
                db.execute("INSERT INTO scenes VALUES(?,?,?,?,?)",
                           (scene_id, payload.video_id, "Автоматический подсчёт", automatic, now()))
                scene = one(db, "scenes", scene_id)
            else:
                scene = dict(scene)
        else:
            raise HTTPException(422, "Выберите видео или сцену")
        # Snapshot the scene so later editor changes do not rewrite old analyses.
        project = Path(__file__).resolve().parents[1]
        def model_path(value):
            raw = Path(value)
            if raw.is_absolute():
                path = raw.resolve()
            elif raw.parts[:2] == ("data", "models"):
                path = (project / raw).resolve()
            elif raw.parts[:1] == ("models",):
                path = (project / "data" / raw).resolve()
            else:
                raise HTTPException(422, "Вес модели должен быть в data/models/*.pt")
            roots = ((project / "data" / "models").resolve(), (ROOT / "models").resolve())
            if path.suffix.lower() != ".pt" or not any(path.is_relative_to(root) for root in roots):
                raise HTTPException(422, "Вес модели должен быть в data/models/*.pt")
            return portable_model_path(path)

        primary = model_path(payload.model or (
            "data/models/yolo11s.pt" if payload.profile == "fast" and
            payload.tracker == "supervision-bytetrack" else "data/models/yolo11m.pt"))
        if payload.tracker != "supervision-bytetrack":
            model = primary
        elif payload.profile == "ensemble":
            secondary = model_path(payload.model_secondary or "data/models/yolov10m.pt")
            if secondary == primary:
                raise HTTPException(422, "Для ансамбля выберите две разные модели")
            model = json.dumps({"models": [primary, secondary], "fusion": "wbf"})
        else:
            model = json.dumps({"pipeline": "fast", "models": [primary],
                                "sample_fps": payload.sample_fps, "imgsz": 640})
        run_geometry = json.loads(scene["geometry"])
        from .cartogram import reference_topology

        example_layout = reference_topology(one(db, "videos", scene["video_id"])["name"])
        run_geometry["topology"] = example_layout if payload.topology == "auto" and example_layout else payload.topology
        run_geometry["topology_source"] = "reference" if payload.topology == "auto" and example_layout else "user"
        db.execute("""INSERT INTO runs(id,scene_id,video_id,geometry,status,model,tracker,created_at)
                      VALUES(?,?,?,?,?,?,?,?)""",
                   (identity, scene["id"], scene["video_id"], json.dumps(run_geometry),
                    "queued", model, payload.tracker, now()))
        return one(db, "runs", identity)


@app.get("/api/runs", tags=["Анализ"], summary="Получить список анализов")
def runs(video_id: str | None = None):
    with connection() as db:
        result = rows(db, "SELECT * FROM runs WHERE video_id=? ORDER BY created_at DESC", (video_id,)) if video_id else rows(db, "SELECT * FROM runs ORDER BY created_at DESC")
    return [run_state(item) for item in result]


def run_state(item):
    """Report a dead worker without rewriting the saved run or its results."""
    if item["status"] == "running" and item["heartbeat"]:
        heartbeat = datetime.fromisoformat(item["heartbeat"])
        if heartbeat < datetime.now(timezone.utc) - timedelta(minutes=5):
            return {**item, "status": "stalled",
                    "message": "Worker не отвечает более пяти минут. Перезапустите Docker и запустите анализ снова."}
        if item["started_at"] and item["progress"] and item["progress"] >= 0.05:
            elapsed = max(0, (datetime.now(timezone.utc) -
                              datetime.fromisoformat(item["started_at"])).total_seconds())
            item["eta_s"] = round(elapsed * (1 - item["progress"]) / item["progress"])
    return item


@app.get("/api/runs/{run_id}", tags=["Анализ"], summary="Получить статус анализа")
def run_detail(run_id: str):
    with connection() as db:
        item = run_state(one(db, "runs", run_id))
        item["geometry"] = json.loads(item["geometry"])
        return item


def change_topology(run_id: str, payload: TopologyIn):
    with connection() as db:
        run = one(db, "runs", run_id)
        geometry = json.loads(run["geometry"])
        geometry["topology"] = payload.topology
        geometry["topology_source"] = "user"
        db.execute("UPDATE runs SET geometry=? WHERE id=?", (json.dumps(geometry), run_id))
        return {"run_id": run_id, "topology": payload.topology}


def cartogram(run_id: str):
    return cartogram_data(aggregate_data(run_id))


def cartogram_data(result):
    from .cartogram import build_cartogram

    track_file = ROOT / "runs" / result["run_id"] / "tracks.json"
    tracks = json.loads(track_file.read_text(encoding="utf-8")) if track_file.is_file() else {}
    return build_cartogram(result["events"], tracks, result["geometry"])


def detections(run_id: str, time_s: float = Query(0, ge=0)):
    """Return tracked boxes at playback time from the finished run, not a new inference."""
    with connection() as db:
        run = one(db, "runs", run_id)
        video = one(db, "videos", run["video_id"])
    folder = ROOT / "runs" / run_id
    index_path = folder / "detections-index.json"
    trace_path = folder / "detections.jsonl"
    if not index_path.is_file() or not trace_path.is_file():
        return {"time_s": time_s, "objects": []}
    index = json.loads(index_path.read_text(encoding="utf-8"))
    seconds = sorted(int(key) for key in index)
    position = bisect_right(seconds, int(time_s)) - 1
    if position < 0:
        return {"time_s": time_s, "objects": []}
    latest = None
    with trace_path.open("rb") as traces:
        traces.seek(index[str(seconds[position])])
        for line in traces:
            frame = json.loads(line)
            if frame["t"] > time_s:
                break
            latest = frame
    # Fast runs may automatically sample down to 2 fps, so retain the last box
    # briefly through the next sampled frame during playback.
    fast_run = run["model"].startswith("{") and json.loads(run["model"]).get("pipeline") == "fast"
    max_age = 0.75 if fast_run else max(0.5, 2 / (video["fps"] or 25))
    if latest is None or time_s - latest["t"] > max_age:
        return {"time_s": time_s, "objects": []}
    return latest


def cancel(run_id: str):
    with connection() as db:
        one(db, "runs", run_id)
        db.execute("UPDATE runs SET cancelled=1,status=CASE WHEN status IN ('queued','running') "
                   "THEN 'cancelled' ELSE status END WHERE id=?", (run_id,))
        return one(db, "runs", run_id)


def events(run_id: str, status: str | None = None):
    with connection() as db:
        one(db, "runs", run_id)
        values = rows(db, "SELECT * FROM events WHERE run_id=? AND status=? ORDER BY time_s", (run_id, status)) if status else rows(db, "SELECT * FROM events WHERE run_id=? ORDER BY time_s", (run_id,))
        return [parse_event(item) for item in values]


def validate_route(run_id, item, db):
    run = one(db, "runs", run_id)
    geometry = json.loads(run["geometry"])
    if geometry.get("mode") == "automatic":
        if (item.from_direction, item.to_direction) != ("В кадре", "В кадре"):
            raise HTTPException(422, "В автоматическом режиме направление: В кадре")
    else:
        if item.from_direction not in {line["name"] for line in geometry["entries"]}:
            raise HTTPException(422, "Неизвестное направление въезда")
        if item.to_direction not in {line["name"] for line in geometry["exits"]}:
            raise HTTPException(422, "Неизвестное направление выезда")
    video = one(db, "videos", run["video_id"])
    if item.time_s > video["duration"] + 1:
        raise HTTPException(422, "Время вне видео")


@app.post("/api/runs/{run_id}/events", tags=["Результаты"], summary="Добавить событие")
def add_event(run_id: str, item: EventIn):
    identity = str(uuid4())
    with connection() as db:
        validate_route(run_id, item, db)
        db.execute("""INSERT INTO events(id,run_id,category,from_direction,to_direction,time_s,start_s,
                      status,origin,details,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
                   (identity, run_id, item.category, item.from_direction, item.to_direction,
                    item.time_s, item.start_s, item.status, "manual", "{}", now(), now()))
        row = parse_event(one(db, "events", identity))
        db.execute("INSERT INTO revisions(event_id,new_value,reason,created_at) VALUES(?,?,?,?)",
                   (identity, json.dumps(row), item.reason, now()))
        return row


@app.patch("/api/runs/{run_id}/events", tags=["Результаты"], summary="Изменить событие")
def edit_event(run_id: str, item: EventEdit):
    with connection() as db:
        old = parse_event(one(db, "events", item.id))
        if old["run_id"] != run_id:
            raise HTTPException(404, "Событие не относится к указанному анализу")
        validate_route(run_id, item, db)
        updated = db.execute("""UPDATE events SET category=?,from_direction=?,to_direction=?,time_s=?,
                        start_s=?,status=?,version=version+1,updated_at=? WHERE id=? AND version=?""",
                     (item.category, item.from_direction, item.to_direction, item.time_s,
                      item.start_s, item.status, now(), item.id, item.version))
        if updated.rowcount != 1:
            raise HTTPException(409, "Запись изменена в другой вкладке. Обновите список.")
        new = parse_event(one(db, "events", item.id))
        db.execute("INSERT INTO revisions(event_id,old_value,new_value,reason,created_at) VALUES(?,?,?,?,?)",
                   (item.id, json.dumps(old), json.dumps(new), item.reason, now()))
        return new


def aggregate_data(run_id, mode="effective", interval=900):
    if interval <= 0 or interval > 86400:
        raise HTTPException(422, "Некорректный интервал")
    with connection() as db:
        run = one(db, "runs", run_id)
        video = one(db, "videos", run["video_id"])
        scene = one(db, "scenes", run["scene_id"])
        all_events = [parse_event(row) for row in rows(db, "SELECT * FROM events WHERE run_id=? ORDER BY time_s", (run_id,))]
    if mode == "effective":
        selected = [event for event in all_events if event["status"] in ("confirmed", "auto_accept")]
    elif mode == "confirmed":
        selected = [event for event in all_events if event["status"] == "confirmed"]
    elif mode == "candidates":
        selected = [event for event in all_events if event["status"] != "discarded"]
    else:
        raise HTTPException(422, "Неизвестный режим агрегирования")
    matrix = Counter((e["from_direction"], e["to_direction"], e["category"]) for e in selected)
    bins = Counter((int(e["time_s"] // interval), e["category"]) for e in selected)
    categories = Counter(event["category"] for event in selected)
    return {"run_id": run_id, "mode": mode, "interval_s": interval,
            "total": len(selected), "statuses": dict(Counter(e["status"] for e in all_events)),
            "category_totals": [{"category": category, "count": count}
                                for category, count in sorted(categories.items())],
            "transport_total": sum(count for category, count in categories.items()
                                   if category != "pedestrian"),
            "pedestrian_total": categories.get("pedestrian", 0),
            "matrix": [{"from": src, "to": dst, "category": cat, "count": n}
                       for (src, dst, cat), n in sorted(matrix.items())],
            "intervals": [{"start_s": idx * interval, "end_s": (idx + 1) * interval,
                          "category": cat, "count": n}
                          for (idx, cat), n in sorted(bins.items())], "events": selected,
            "geometry": json.loads(run["geometry"]),
            "video": {key: video[key] for key in
                      ("id", "name", "width", "height", "duration", "fps", "created_at")},
            "scene": {"id": scene["id"], "name": scene["name"]}}


def aggregate(run_id: str, mode: str = "effective", interval: int = 900):
    return aggregate_data(run_id, mode, interval)


@app.patch(
    "/api/runs/{run_id}",
    tags=["Анализ"],
    summary="Изменить анализ",
    description="Один endpoint для остановки анализа или смены вида перекрёстка.",
)
def update_run(run_id: str, payload: RunUpdate):
    if payload.action == "cancel":
        return cancel(run_id)
    return change_topology(run_id, TopologyIn(topology=payload.topology))


@app.get(
    "/api/runs/{run_id}/results",
    tags=["Результаты"],
    summary="Получить результаты анализа",
    description=(
        "По умолчанию возвращает сводку, события и картограмму одним ответом. "
        "Параметр `view` позволяет запросить только один раздел или рамки объектов для кадра."
    ),
)
def results(
    run_id: str,
    view: Literal["all", "summary", "events", "cartogram", "detections"] = "all",
    mode: Literal["effective", "confirmed", "candidates"] = "effective",
    interval: int = Query(900, gt=0, le=86400),
    time_s: float = Query(0, ge=0),
    status: str | None = None,
):
    if view == "detections":
        return detections(run_id, time_s)
    if view == "events":
        return events(run_id, status)
    summary = aggregate_data(run_id, mode, interval)
    if view == "summary":
        return summary
    diagram = cartogram_data(summary)
    if view == "cartogram":
        return diagram
    return {"summary": summary, "events": events(run_id, status), "cartogram": diagram}


def safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


CATEGORY_LABELS = {
    "pedestrian": "Пешеходы",
    "bicycle": "Вело",
    "car": "Авто",
    "motorcycle": "Мото",
    "truck": "Груз.",
    "bus": "ОТ",
}
CATEGORY_ORDER = ("pedestrian", "bicycle", "car", "motorcycle", "truck", "bus")
MODE_LABELS = {"effective": "Учтённые", "confirmed": "Только подтверждённые",
               "candidates": "Все кандидаты"}
TOPOLOGY_LABELS = {"three_way": "Трёхлучевой", "four_way": "Четырёхлучевой",
                   "unknown": "Не определён"}
REPORT_NAVY = "0B1F3A"
REPORT_BLUE = "1E5AA8"
REPORT_PALE = "EAF2FB"
REPORT_BORDER = "C8D5E6"


def _report_video_parts(name):
    """Extract period and recording/circle number used by the supplied examples."""
    stem = Path(name).stem
    match = re.search(r"(?:^|[_\s-])(утро|вечер|выходной)(?:[_\s-]+)(\d+)",
                      stem, re.IGNORECASE)
    if not match:
        return "", ""
    return match.group(1).lower(), int(match.group(2))


def _style_table_sheet(page, freeze="A2", filter_table=True):
    """Apply restrained dark-blue report styling to a rectangular data sheet."""
    if page.max_row < 1 or page.max_column < 1:
        return
    dark = PatternFill("solid", fgColor=REPORT_NAVY)
    line = Side(style="thin", color=REPORT_BORDER)
    for cell in page[1]:
        cell.fill = dark
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=line)
    for row in page.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="172033")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=Side(style="hair", color="E7EDF5"))
    page.freeze_panes = freeze
    if filter_table and page.max_row >= 2:
        page.auto_filter.ref = page.dimensions
    page.sheet_view.showGridLines = False
    page.row_dimensions[1].height = 28
    for column in range(1, page.max_column + 1):
        values = [str(page.cell(row, column).value or "") for row in range(1, page.max_row + 1)]
        page.column_dimensions[get_column_letter(column)].width = min(42, max(10, max(map(len, values)) + 2))


def _flow_groups(result, diagram):
    """Create one reconciled set of report directions without double-counting."""
    groups = defaultdict(Counter)
    labels = {}
    if result["geometry"].get("mode") == "automatic" and diagram["topology"] != "unknown":
        for row in diagram["direction_matrix"]:
            key = (row["direction"], row["from"], row["to"])
            labels[key] = f"Направление {row['direction']}\n{row['from']} → {row['to']}"
            groups[key][row["category"]] += row["count"]
        if diagram["unresolved"]:
            key = (10_000, "Не определено", "Не определено")
            labels[key] = "Направление не определено"
            for row in diagram["unresolved_by_category"]:
                groups[key][row["category"]] += row["count"]
    else:
        for row in result["matrix"]:
            key = (0, row["from"], row["to"])
            groups[key][row["category"]] += row["count"]
        ordered = sorted(groups, key=lambda key: (key[1], key[2]))
        groups = {key: groups[key] for key in ordered}
        labels = {key: f"Направление {index}\n{key[1]} → {key[2]}"
                  for index, key in enumerate(ordered, 1)}
    return [(key, labels[key], groups[key]) for key in groups]


def _add_count_sheet(book, result, diagram):
    """Add a wide two-level count matrix matching the supplied workbook convention."""
    page = book.create_sheet("Подсчёты")
    period, recording = _report_video_parts(result["video"]["name"])
    metadata = (("Видео", result["video"]["name"]), ("Сцена", result["scene"]["name"]),
                ("Период", period), ("Видео / круг", recording))
    for column, (heading, value) in enumerate(metadata, 1):
        page.merge_cells(start_row=1, start_column=column, end_row=2, end_column=column)
        page.cell(1, column, safe_cell(heading))
        page.cell(3, column, safe_cell(value))

    groups = _flow_groups(result, diagram)
    observed = {row["category"] for row in result["category_totals"]}
    categories = [category for category in CATEGORY_ORDER if category in observed]
    categories.extend(sorted(observed - set(categories)))
    if not categories:
        categories = list(CATEGORY_ORDER)
    column = 5
    for _, label, counts in groups:
        start = column
        end = column + len(categories) - 1
        if end > start:
            page.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        page.cell(1, start, safe_cell(label))
        for category in categories:
            page.cell(2, column, safe_cell(CATEGORY_LABELS.get(category, category)))
            page.cell(3, column, counts.get(category, 0))
            column += 1
    total_start = column
    total_end = column + len(categories) - 1
    if total_end > total_start:
        page.merge_cells(start_row=1, start_column=total_start, end_row=1, end_column=total_end)
    page.cell(1, total_start, "Итого")
    totals = {row["category"]: row["count"] for row in result["category_totals"]}
    for category in categories:
        page.cell(2, column, safe_cell(CATEGORY_LABELS.get(category, category)))
        page.cell(3, column, totals.get(category, 0))
        column += 1

    dark = PatternFill("solid", fgColor=REPORT_NAVY)
    blue = PatternFill("solid", fgColor=REPORT_BLUE)
    pale = PatternFill("solid", fgColor=REPORT_PALE)
    line = Side(style="thin", color="FFFFFF")
    for row in (1, 2):
        for cell in page[row]:
            cell.fill = dark if row == 1 else blue
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=line, right=line, bottom=line)
    for cell in page[3]:
        cell.fill = pale
        cell.font = Font(name="Arial", size=10, color="172033")
        cell.alignment = Alignment(horizontal="right" if cell.column >= 5 else "left",
                                   vertical="center")
        cell.border = Border(bottom=Side(style="thin", color=REPORT_BORDER))
    page.freeze_panes = "E3"
    page.sheet_view.showGridLines = False
    page.row_dimensions[1].height = 38
    page.row_dimensions[2].height = 26
    page.column_dimensions["A"].width = 30
    page.column_dimensions["B"].width = 24
    page.column_dimensions["C"].width = 14
    page.column_dimensions["D"].width = 14
    for index in range(5, page.max_column + 1):
        page.column_dimensions[get_column_letter(index)].width = 11
    page.print_title_rows = "1:2"
    page.print_title_cols = "A:D"
    page.page_setup.orientation = "landscape"
    page.page_setup.scale = 80
    return page


def _add_summary_sheet(book, result, diagram):
    page = book.create_sheet("Отчёт", 0)
    page["A2"] = "Отчёт по транспортным и пешеходным потокам"
    page["A2"].font = Font(name="Arial", size=15, bold=True, color=REPORT_NAVY)
    page.merge_cells("A2:F2")
    page["A3"] = "Результаты выбранного запуска анализа"
    page["A3"].font = Font(name="Arial", size=10, italic=True, color="52647A")
    page.merge_cells("A3:F3")
    page["A4"].fill = PatternFill("solid", fgColor=REPORT_BLUE)
    page.merge_cells("A4:F4")

    details = [
        ("Видео", result["video"]["name"]),
        ("Сцена", result["scene"]["name"]),
        ("Режим данных", MODE_LABELS.get(result["mode"], result["mode"])),
        ("Интервал", f"{result['interval_s']} с"),
        ("Тип узла", TOPOLOGY_LABELS.get(diagram["topology"], diagram["topology"])),
    ]
    for row, (label, value) in enumerate(details, 6):
        page.cell(row, 1, label)
        page.cell(row, 2, safe_cell(value))
        page.cell(row, 1).font = Font(name="Arial", size=10, bold=True, color=REPORT_NAVY)
        page.cell(row, 2).font = Font(name="Arial", size=10, color="172033")

    metrics = (("Всего учтено", result["total"]),
               ("Транспорт", result["transport_total"]),
               ("Пешеходы", result["pedestrian_total"]),
               ("Направление определено", diagram["routed"]),
               ("Направление не определено", diagram["unresolved"]))
    for row, (label, value) in enumerate(metrics, 6):
        page.cell(row, 4, label)
        page.cell(row, 5, value)
        page.cell(row, 4).fill = PatternFill("solid", fgColor=REPORT_PALE)
        page.cell(row, 5).fill = PatternFill("solid", fgColor=REPORT_PALE)
        page.cell(row, 4).font = Font(name="Arial", size=10, bold=True, color=REPORT_NAVY)
        page.cell(row, 5).font = Font(name="Arial", size=11, bold=True, color=REPORT_BLUE)
        page.cell(row, 5).alignment = Alignment(horizontal="right")

    page["A13"] = "По категориям"
    page["A13"].fill = PatternFill("solid", fgColor=REPORT_NAVY)
    page["B13"] = "Количество"
    page["B13"].fill = PatternFill("solid", fgColor=REPORT_NAVY)
    for cell in page[13][:2]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    for row, item in enumerate(result["category_totals"], 14):
        page.cell(row, 1, safe_cell(CATEGORY_LABELS.get(item["category"], item["category"])))
        page.cell(row, 2, item["count"])

    page["D13"] = "Покрытие картограммы"
    page["D13"].fill = PatternFill("solid", fgColor=REPORT_NAVY)
    page["E13"] = "Значение"
    page["E13"].fill = PatternFill("solid", fgColor=REPORT_NAVY)
    for cell in (page["D13"], page["E13"]):
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    coverage = (("Доля определённых", diagram["coverage"]),
                ("По движению в кадре", diagram["motion_estimated"]),
                ("По именованным линиям", diagram["user_directions"]))
    for row, (label, value) in enumerate(coverage, 14):
        page.cell(row, 4, label)
        page.cell(row, 5, value)
    page["E14"].number_format = "0.0%"

    note_row = max(20, 15 + len(result["category_totals"]))
    page.cell(note_row, 1, "Примечание")
    page.cell(note_row, 1).font = Font(name="Arial", size=10, bold=True, color=REPORT_NAVY)
    page.cell(note_row + 1, 1,
              "Номера направлений в листе «Подсчёты» относятся к сформированной схеме. "
              "Для сопоставления с нумерацией приложенной картограммы нужна явная привязка линий камеры.")
    page.merge_cells(start_row=note_row + 1, start_column=1, end_row=note_row + 2, end_column=6)
    page.cell(note_row + 1, 1).alignment = Alignment(wrap_text=True, vertical="top")
    page.cell(note_row + 1, 1).font = Font(name="Arial", size=10, color="52647A")
    page.sheet_view.showGridLines = False
    page.column_dimensions["A"].width = 27
    page.column_dimensions["B"].width = 32
    page.column_dimensions["C"].width = 4
    page.column_dimensions["D"].width = 29
    page.column_dimensions["E"].width = 16
    page.column_dimensions["F"].width = 4
    page.print_area = f"A1:F{note_row + 2}"
    page.page_setup.orientation = "landscape"
    page.page_setup.fitToWidth = 1
    page.page_setup.fitToHeight = 1
    page.sheet_properties.pageSetUpPr.fitToPage = True
    page.page_margins.left = 0.25
    page.page_margins.right = 0.25
    page.page_margins.top = 0.35
    page.page_margins.bottom = 0.35
    page.print_options.horizontalCentered = True
    return page


@app.get("/api/runs/{run_id}/export/{format}", tags=["Результаты"], summary="Скачать отчёт")
def export(run_id: str, format: str, mode: str = "effective", interval: int = 900):
    result = aggregate_data(run_id, mode, interval)
    if format == "json":
        return Response(json.dumps({**result, "cartogram": cartogram_data(result)}, ensure_ascii=False), media_type="application/json",
                        headers={"Content-Disposition": f'attachment; filename="{run_id}.json"'})
    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(("from", "to", "category", "count"))
        for row in result["matrix"]:
            writer.writerow([safe_cell(row[key]) for key in ("from", "to", "category", "count")])
        return Response("\ufeff" + output.getvalue(), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="{run_id}.csv"'})
    if format == "xlsx":
        diagram = cartogram_data(result)
        book = Workbook()
        book.remove(book.active)
        _add_summary_sheet(book, result, diagram)
        _add_count_sheet(book, result, diagram)
        for name, headings, records in (
            ("Summary", ("metric", "value"),
             [("video", result["video"]["name"]), ("scene", result["scene"]["name"]),
              ("total", result["total"]), ("transport", result["transport_total"]),
              ("pedestrians", result["pedestrian_total"]), ("mode", mode),
              ("interval_s", interval)]),
            ("Flow Matrix", ("from", "to", "category", "count"),
              [[row[k] for k in ("from", "to", "category", "count")] for row in result["matrix"]]),
            ("Cartogram estimate", ("from (schematic)", "to (schematic)", "category", "count"),
             [[row[k] for k in ("from", "to", "category", "count")] for row in diagram["matrix"]]),
            ("Cartogram coverage", ("metric", "value"),
             [("topology", diagram["topology"]), ("routed estimate", diagram["routed"]),
              ("motion estimated", diagram["motion_estimated"]),
              ("user directions", diagram["user_directions"]),
              ("unresolved", diagram["unresolved"]), ("all counted", diagram["total"])]),
            ("Intervals", ("start_s", "end_s", "category", "count"),
             [[row[k] for k in ("start_s", "end_s", "category", "count")] for row in result["intervals"]]),
            ("Events", ("id", "time_s", "from", "to", "category", "status", "origin",
                        "tracker_id", "confidence"),
             [[e.get(k) for k in ("id", "time_s", "from_direction", "to_direction", "category",
                                  "status", "origin", "tracker_id", "confidence")]
              for e in result["events"]]),
            ("Review", ("status", "count"), list(result["statuses"].items())),
        ):
            page = book.create_sheet(name)
            page.append(headings)
            for record in records:
                page.append([safe_cell(value) for value in record])
            _style_table_sheet(page)
        book.active = 0
        data = io.BytesIO()
        book.save(data)
        return Response(data.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{run_id}.xlsx"'})
    raise HTTPException(422, "Форматы: xlsx, csv, json")


@app.post("/api/videos/{video_id}/ground-truth", include_in_schema=False)
async def ground_truth(video_id: str, file: UploadFile = File(...)):
    raw = await file.read(10_000_001)
    if len(raw) > 10_000_000:
        raise HTTPException(413, "Эталон больше 10 МБ")
    try:
        extension = Path(file.filename or "").suffix.lower()
        if extension == ".json":
            payload, format_ = json.loads(raw), "json"
        elif extension == ".csv":
            payload, format_ = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))), "csv"
        elif extension == ".xlsx":
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            values = book.active.values
            header = [str(value) if value is not None else "" for value in next(values)]
            payload = [dict(zip(header, row)) for row in values]
            format_ = "xlsx"
        else:
            raise ValueError("Поддерживаются JSON, CSV, XLSX")
        if not isinstance(payload, (list, dict)):
            raise ValueError("Нужен массив событий или таблица")
    except (UnicodeError, ValueError) as exc:
        raise HTTPException(422, f"Некорректный эталон: {exc}") from exc
    identity = str(uuid4())
    with connection() as db:
        one(db, "videos", video_id)
        db.execute("INSERT INTO ground_truth VALUES(?,?,?,?,?,?)",
                   (identity, video_id, file.filename, format_, json.dumps(payload), now()))
    return {"id": identity, "format": format_}


@app.get("/api/runs/{run_id}/evaluate/{reference_id}", include_in_schema=False)
def evaluate(run_id: str, reference_id: str):
    with connection() as db:
        run = one(db, "runs", run_id)
        gt = one(db, "ground_truth", reference_id)
    if gt["video_id"] != run["video_id"]:
        raise HTTPException(422, "Эталон относится к другому видео")
    reference = json.loads(gt["payload"])
    if isinstance(reference, dict):
        reference = reference.get("events", reference.get("matrix", []))
    if not isinstance(reference, list):
        raise HTTPException(422, "Некорректный формат эталона")
    expected = Counter()
    for row in reference:
        src = row.get("from_direction", row.get("from"))
        dst = row.get("to_direction", row.get("to"))
        cat = row.get("category")
        if src is None or dst is None or cat is None:
            raise HTTPException(422, "Нужны столбцы from/to/category")
        expected[(src, dst, cat)] += int(row.get("count", 1))
    actual = Counter({(row["from"], row["to"], row["category"]): row["count"]
                      for row in aggregate_data(run_id, "candidates")["matrix"]})
    keys = expected.keys() | actual.keys()
    absolute_error = sum(abs(expected[key] - actual[key]) for key in keys)
    gt_total = sum(expected.values())
    return {"ground_truth": gt_total, "predicted": sum(actual.values()),
            "absolute_error": absolute_error,
            "wape": absolute_error / gt_total if gt_total else None,
            "relative_error": None if gt_total == 0 else (sum(actual.values()) - gt_total) / gt_total,
            "rows": [{"from": a, "to": b, "category": c, "ground_truth": expected[a,b,c],
                      "predicted": actual[a,b,c]} for a,b,c in sorted(keys)]}
