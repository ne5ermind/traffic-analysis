"""Opt-in end-to-end check against running API/Redis/worker with a real video.

Leaves the project and exports available for inspection; no synthetic counts.
Requires development dependencies (httpx).
"""

import argparse
import json
import time
from pathlib import Path
import httpx
from openpyxl import load_workbook


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--video", type=Path, required=True)
    parser.add_argument("--base-url", default="http://localhost:3000/api")
    parser.add_argument("--output", type=Path, default=Path("data/http-smoke"))
    parser.add_argument("--timeout", type=int, default=1800)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    with httpx.Client(base_url=args.base_url.rstrip("/"), timeout=120) as client:

        def request(method, path, **kwargs):
            response = client.request(method, path, **kwargs)
            response.raise_for_status()
            return response.json()

        project = request(
            "POST",
            "/projects",
            json={"name": "Проверка Docker · реальное видео", "description": "Техническая end-to-end проверка; без ручного эталона точности."},
        )
        root = f"/projects/{project['id']}"
        with args.video.open("rb") as video:
            request("POST", root + "/video", files={"file": (args.video.name, video, "video/mp4")})
        request("POST", root + "/analysis", json={"profile": "balanced"})
        deadline = time.monotonic() + args.timeout
        last = None
        while time.monotonic() < deadline:
            status = request("GET", root + "/analysis")
            marker = (status["status"], status.get("stage"), status.get("percent"))
            if marker != last:
                print(json.dumps(status, ensure_ascii=False), flush=True)
                last = marker
            if status["status"] == "completed":
                break
            if status["status"] in ("failed", "cancelled"):
                raise RuntimeError(status.get("error") or status["status"])
            time.sleep(2)
        else:
            raise TimeoutError(f"Обработка продолжается: {project['id']}")
        project = request("GET", root)
        for ext in ("xlsx", "csv", "json", "svg"):
            with client.stream("GET", root + "/export/" + ext) as response:
                response.raise_for_status()
                with (args.output / f"result.{ext}").open("wb") as output:
                    for chunk in response.iter_bytes():
                        output.write(chunk)
        workbook = load_workbook(args.output / "result.xlsx", read_only=True)
        assert workbook.sheetnames == ["Summary", "Movements", "Time intervals", "Tracks", "Metadata"]
        workbook.close()
        response = client.get(root + "/media/playback", headers={"Range": "bytes=0-99"})
        assert response.status_code == 206 and len(response.content) == 100
        (args.output / "project.json").write_text(json.dumps(project, ensure_ascii=False, indent=2))
        print(
            json.dumps(
                {
                    "project_id": project["id"],
                    "counts": project["result"]["counts"],
                    "diagnostics": project["result"]["diagnostics"],
                    "exports": str(args.output),
                },
                ensure_ascii=False,
                indent=2,
            )
        )


if __name__ == "__main__":
    main()
